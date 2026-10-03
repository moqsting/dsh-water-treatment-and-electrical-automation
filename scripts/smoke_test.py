# -*- coding: utf-8 -*-
r"""
smoke_test.py —— 整合包冒烟测试

在无 AutoCAD、无真实图纸的机器上验证核心链路：
1. 生成测试 DXF（含设备块+位号属性、仪表文本）→ dxf_parse 提取
2. 生成招标清单与两家报价 xlsx → quote_tools 规整/比对/测算/差异
3. cad_env 探测（应给出"未找到+指引"而非报错）

用法：integration-pack\scripts\runpy.cmd integration-pack\scripts\smoke_test.py
"""
import subprocess
import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

import pandas as pd

PACK = Path(__file__).resolve().parent.parent
WORK = PACK / "reports" / "smoke"
sys.path.insert(0, str(PACK / "pydeps"))


def step(name):
    print(f"\n=== {name} ===")


def run_py(script, *args):
    cmd = [sys.executable, str(PACK / "scripts" / script), *args]
    r = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace")
    if r.returncode != 0:
        print("  失败：", (r.stderr or r.stdout)[:400])
        return False
    print("  OK:", (r.stdout or "").strip().splitlines()[-1] if (r.stdout or "").strip() else "")
    return True


def expect_fail(script, expect_substrs, *args, label=""):
    """运行脚本并断言：退出码非 0 且输出包含全部期望子串（数据契约错误场景）。"""
    subs = [expect_substrs] if isinstance(expect_substrs, str) else expect_substrs
    cmd = [sys.executable, str(PACK / "scripts" / script), *args]
    r = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace")
    out = (r.stderr or "") + (r.stdout or "")
    missing = [s for s in subs if s not in out]
    ok = r.returncode != 0 and not missing
    if ok:
        print(f"  PASS(契约): {label} → 退出码 {r.returncode}，错误信息含全部期望关键词")
    else:
        print(f"  FAIL(契约): {label} → 退出码 {r.returncode}，缺失关键词 {missing}，实际输出：{out[:300]}")
    return ok


def make_test_dxf():
    import ezdxf
    doc = ezdxf.new("R2018")
    msp = doc.modelspace()
    doc.blocks.new(name="PUMP")
    doc.blocks.new(name="TANK")
    for i, (blk, tag) in enumerate([("PUMP", "P-101"), ("PUMP", "P-102"), ("TANK", "V-201")]):
        ins = msp.add_blockref(blk, (i * 20, 0))
        ins.add_attrib("TAG", tag, (0, 5))
    msp.add_text("FIT-101 INLET FLOW").set_placement((0, 15))
    msp.add_text("LIT-201 TANK LEVEL").set_placement((20, 15))
    out = WORK / "测试PID.dxf"
    doc.saveas(out)
    return out


def make_quote_xlsx():
    import pandas as pd
    tender = pd.DataFrame({
        "设备名称": ["潜水排污泵", "电磁流量计", "PAC加药装置"],
        "规格型号": ["Q=100m3/h H=15m", "DN200", "500L/h"],
        "材质": ["铸铁", "衬四氟", "PE"],
        "单位": ["台", "台", "套"],
        "数量": [2, 2, 1],
    })
    q1 = tender.copy()
    q1["单价(含税)"] = [18500, 3200, 45000]
    q1["品牌"] = ["A牌", "B牌", "C牌"]
    q2 = tender.copy()
    q2["单价(含税)"] = [19800, 3000, 52000]
    q2["品牌"] = ["X牌", "B牌", "Y牌"]
    t = WORK / "招标清单.xlsx"
    a = WORK / "报价A.xlsx"
    b = WORK / "报价B.xlsx"
    tender.to_excel(t, index=False)
    q1.to_excel(a, index=False)
    q2.to_excel(b, index=False)
    return t, a, b


def main():
    WORK.mkdir(parents=True, exist_ok=True)

    step("cad_env 探测（预期：未找到+指引）")
    run_py("cad_env.py")

    step("生成测试 DXF")
    dxf = make_test_dxf()
    print("  已生成:", dxf)

    step("dxf_parse 解析")
    ok = run_py("dxf_parse.py", str(dxf), "--out", str(WORK / "提取清单.xlsx"))
    if ok:
        import openpyxl
        wb = openpyxl.load_workbook(WORK / "提取清单.xlsx")
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        print(f"  提取 {len(rows) - 1} 条：")
        for r in rows[1:]:
            print("   ", r[0], r[1], r[2], "置信度", r[7], r[8])

    step("生成测试报价表")
    t, a, b = make_quote_xlsx()
    print("  已生成:", t.name, a.name, b.name)

    step("quote_tools normalize")
    run_py("quote_tools.py", "normalize", str(t), "--out", str(WORK / "设备表.xlsx"))

    step("quote_tools compare")
    run_py("quote_tools.py", "compare", str(a), str(b), "--out", str(WORK / "比对表.xlsx"))

    step("quote_tools cost")
    run_py("quote_tools.py", "cost", str(a),
           "--tax", "13", "--freight", "5000", "--install", "3", "--admin", "5",
           "--profit", "8", "--out", str(WORK / "测算表.xlsx"))

    step("quote_tools diff")
    run_py("quote_tools.py", "diff", str(t), str(a), "--out", str(WORK / "差异报告.xlsx"))

    # ── 数据契约场景（工具间串联约束）──
    step("数据契约：normalize 输出 → cost（缺'单价'，应报契约错误而非 KeyError）")
    expect_fail(
        "quote_tools.py",
        ["Missing required field: 单价", "Required by: cost", "Available fields", "compare"],
        "cost", str(WORK / "设备表.xlsx"), "--out", str(WORK / "不应生成.xlsx"),
        label="normalize→cost 不可直连",
    )

    step("数据契约：字段类型错误（单价列含文字）")
    bad = WORK / "单价含文字.xlsx"
    pd.DataFrame({"设备名称": ["潜污泵"], "单价": ["abc"]}).to_excel(bad, index=False)
    expect_fail(
        "quote_tools.py",
        ["字段类型错误", "单价", "Required by: cost"],
        "cost", str(bad), "--out", str(WORK / "不应生成2.xlsx"),
        label="单价类型错误",
    )

    step("数据契约：空数据（只有表头无数据行）")
    empty = WORK / "空表.xlsx"
    pd.DataFrame({"单价": pd.Series(dtype="float64")}).to_excel(empty, index=False)
    expect_fail(
        "quote_tools.py",
        ["空数据", "Required by: cost"],
        "cost", str(empty), "--out", str(WORK / "不应生成3.xlsx"),
        label="空数据拒绝",
    )

    step("数据契约：CSV 输入（中文表头，utf-8-sig）")
    csv_in = WORK / "招标清单中文.csv"
    pd.DataFrame({
        "设备名称": ["潜水排污泵", "电磁流量计"],
        "规格型号": ["Q=100m3/h H=15m", "DN200"],
        "数量": [2, 2],
    }).to_csv(csv_in, index=False, encoding="utf-8-sig")
    run_py("quote_tools.py", "normalize", str(csv_in), "--out", str(WORK / "设备表_来自csv.xlsx"))

    step("数据契约：compare 输出 → cost（含单价可直连；同设备多行默认拒绝，需 --allow-duplicates）")
    expect_fail(
        "quote_tools.py",
        ["重复条目", "重复计价", "Required by: cost"],
        "cost", str(WORK / "比对表.xlsx"), "--tax", "13", "--out", str(WORK / "不应生成4.xlsx"),
        label="重复键默认拒绝",
    )
    run_py("quote_tools.py", "cost", str(WORK / "比对表.xlsx"),
           "--tax", "13", "--allow-duplicates", "--out", str(WORK / "测算表_来自比对.xlsx"))

    step("数据契约：中文字段（normalize 中文表头输出含中文内容保真）")
    import openpyxl as _xl
    wb2 = _xl.load_workbook(WORK / "设备表_来自csv.xlsx")
    ws2 = wb2.active
    rows2 = list(ws2.iter_rows(values_only=True))
    print("  中文表头:", rows2[0][:4])
    print("  中文内容:", rows2[1][:4])

    print("\n冒烟测试完成，产物在", WORK)


if __name__ == "__main__":
    main()
