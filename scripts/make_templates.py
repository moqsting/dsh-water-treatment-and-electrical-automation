# -*- coding: utf-8 -*-
"""make_templates.py —— 生成整合包 Excel 模板（投标报价/IO点表/设备清单/电缆清册/调试记录）"""
import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    print("缺少 openpyxl：请先运行 runpy.cmd bootstrap_libs.py")
    sys.exit(1)

TEMPLATES = Path(__file__).resolve().parent.parent / "templates"

HEADER_FILL = PatternFill("solid", fgColor="D9E2F3")
HEADER_FONT = Font(bold=True)


def _write_table(ws, headers, widths=None, start_row=2):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=c, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def make_quote():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "分项报价"
    ws["A1"] = "投标分项报价表（按招标文件格式调整）"
    ws["A1"].font = Font(bold=True, size=13)
    _write_table(ws, ["序号", "设备名称", "规格型号", "品牌", "材质", "单位", "数量",
                      "单价(含税)", "合计(含税)", "货期", "质保", "备注"],
                 [6, 22, 30, 12, 10, 8, 8, 12, 12, 10, 10, 14])
    ws2 = wb.create_sheet("汇总")
    ws2["A1"] = "报价汇总表"
    ws2["A1"].font = Font(bold=True, size=13)
    _write_table(ws2, ["项目", "金额(元)", "说明"], [24, 14, 40])
    for r, item in enumerate(["设备合计", "运费", "安装费", "管理费", "利润", "税金", "含税总价"], 3):
        ws2.cell(row=r, column=1, value=item)
    wb.save(TEMPLATES / "投标报价表模板.xlsx")


def make_io():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "IO点表"
    ws["A1"] = "I/O 点表（PLC）"
    ws["A1"].font = Font(bold=True, size=13)
    _write_table(ws, ["序号", "位号", "名称/描述", "信号类型(DI/DO/AI/AO)", "量程/单位",
                      "信号制(4-20mA等)", "供电", "柜号/模块", "PLC地址", "备注"],
                 [6, 14, 26, 18, 12, 12, 10, 12, 12, 16])
    wb.save(TEMPLATES / "IO点表模板.xlsx")


def make_bom():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "设备清单"
    ws["A1"] = "设备清单（招标/投标用）"
    ws["A1"].font = Font(bold=True, size=13)
    _write_table(ws, ["序号", "设备名称", "规格型号", "材质", "单位", "数量", "功率(kW)",
                      "电压等级", "品牌(投标填写)", "备注"],
                 [6, 22, 30, 10, 8, 8, 10, 10, 14, 16])
    wb.save(TEMPLATES / "设备清单模板.xlsx")


def make_cable():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "电缆清册"
    ws["A1"] = "电缆清册"
    ws["A1"].font = Font(bold=True, size=13)
    _write_table(ws, ["序号", "电缆编号", "起点设备", "终点设备", "型号规格", "长度(m)",
                      "敷设方式", "电压等级", "备注"],
                 [6, 14, 18, 18, 20, 10, 12, 10, 16])
    wb.save(TEMPLATES / "电缆清册模板.xlsx")


def make_debug():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "调试记录"
    ws["A1"] = "设备/回路调试记录"
    ws["A1"].font = Font(bold=True, size=13)
    _write_table(ws, ["日期", "设备/回路", "调试项目", "调试方法", "结果(合格/不合格)",
                      "问题描述", "处理措施", "调试人", "复核人"],
                 [12, 16, 20, 20, 16, 24, 24, 10, 10])
    wb.save(TEMPLATES / "调试记录模板.xlsx")


if __name__ == "__main__":
    TEMPLATES.mkdir(parents=True, exist_ok=True)
    make_quote(); make_io(); make_bom(); make_cable(); make_debug()
    print("模板已生成于", TEMPLATES)
