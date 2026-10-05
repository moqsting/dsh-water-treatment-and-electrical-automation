# -*- coding: utf-8 -*-
"""ui/tests/make_fixtures.py —— 生成集成测试所需夹具（由 run_integration.ps1 调用）。

docx 预览用例需要一份含「预览测试」字样的文档。此前该夹具依赖手工遗留文件，
换台机器或清空 reports 后就缺失，用例即失败；改为测试前按需生成。

不用 `py -3 -c "<含中文的代码>"` 的形式：Windows 下 py launcher 转发命令行参数时
会经系统 ANSI（中文系统为 GBK）往返，中文源码会变成乱码并触发 NameError。
以文件形式传入则 Python 按 UTF-8 读取源码，不受代码页影响。

用法：py -3 ui/tests/make_fixtures.py [pack_root]
"""
import sys
from pathlib import Path


def make_docx(target: Path) -> None:
    """生成含两段中文的 docx（用于 /api/file 的 docx 预览断言）。"""
    from docx import Document

    target.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_paragraph("预览测试段落一")
    doc.add_paragraph("预览测试段落二")
    doc.save(str(target))


def main() -> int:
    if len(sys.argv) > 1:
        pack_root = Path(sys.argv[1]).resolve()
    else:
        pack_root = Path(__file__).resolve().parent.parent.parent  # ui/tests -> 包根

    target = pack_root / "reports" / "ui" / "4i-test.docx"
    try:
        make_docx(target)
    except ImportError as exc:
        print(f"缺少 python-docx（{exc}）：请确认 pydeps 已就绪并已加入 PYTHONPATH")
        return 1
    if not target.is_file():
        print(f"夹具未生成：{target}")
        return 1
    print(f"已生成夹具：{target}（{target.stat().st_size} 字节）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
